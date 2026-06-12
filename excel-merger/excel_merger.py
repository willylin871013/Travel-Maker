# -*- coding: utf-8 -*-
"""Excel 批次合併小工具

把資料夾內所有 .xlsx 檔案合併成一份總表，並自動產生統計摘要。
適用情境:不同機台/班別/日期的報表格式相同,需要定期彙整。

使用方式:
    1. 點「產生示範資料」可以先建立三個假的機台報表來測試
    2. 選擇來源資料夾後點「開始合併」
    3. 合併結果會輸出在同一個資料夾,檔名含時間戳記
"""

import datetime
import random
import threading
from pathlib import Path

try:  # 沒有圖形環境(例如 CI 測試)時仍可使用合併邏輯
    import tkinter as tk
    from tkinter import filedialog, messagebox, ttk
except ImportError:
    tk = None

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

APP_TITLE = "Excel 批次合併小工具 v1.0"
SOURCE_COLUMN = "來源檔案"


# ---------------------------------------------------------------------------
# 核心邏輯(與介面無關,方便之後改成命令列或排程執行)
# ---------------------------------------------------------------------------

def find_excel_files(folder: Path) -> list[Path]:
    """列出資料夾內所有 .xlsx 檔(略過 Excel 暫存檔與先前的合併結果)。"""
    return sorted(
        f for f in folder.glob("*.xlsx")
        if not f.name.startswith("~$") and not f.name.startswith("合併結果_")
    )


def merge_excel_files(files: list[Path], log=print) -> Path:
    """把多個 .xlsx 合併成一份總表,回傳輸出檔路徑。

    以第一個檔案的標題列為準,其餘檔案依標題名稱對齊欄位,
    所以欄位順序不同也能正確合併。
    """
    headers: list[str] = []
    rows: list[list] = []          # 合併後的資料列(含來源檔名)
    per_file_count: dict[str, int] = {}

    for f in files:
        wb = load_workbook(f, data_only=True, read_only=True)
        ws = wb.active
        file_rows = ws.iter_rows(values_only=True)

        file_headers = [str(h) if h is not None else "" for h in next(file_rows, [])]
        if not headers:
            headers = file_headers
        # 依標題名稱對齊到基準欄位;找不到的欄位留空
        index_map = []
        for h in headers:
            index_map.append(file_headers.index(h) if h in file_headers else None)

        count = 0
        for row in file_rows:
            if all(v is None for v in row):
                continue
            aligned = [row[i] if i is not None and i < len(row) else None for i in index_map]
            rows.append(aligned + [f.name])
            count += 1
        wb.close()
        per_file_count[f.name] = count
        log(f"  讀取 {f.name}:{count} 筆資料")

    out = Workbook()

    # --- 工作表 1:合併資料 ---
    ws_data = out.active
    ws_data.title = "合併資料"
    all_headers = headers + [SOURCE_COLUMN]
    ws_data.append(all_headers)
    for cell in ws_data[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="4472C4")
    for row in rows:
        ws_data.append(row)
    ws_data.freeze_panes = "A2"
    ws_data.auto_filter.ref = f"A1:{get_column_letter(len(all_headers))}{len(rows) + 1}"

    # --- 工作表 2:統計摘要 ---
    ws_stat = out.create_sheet("統計摘要")
    ws_stat.append(["統計項目", "數值"])
    ws_stat.append(["合併檔案數", len(files)])
    ws_stat.append(["總資料筆數", len(rows)])
    ws_stat.append([])
    ws_stat.append(["來源檔案", "資料筆數"])
    chart_start = ws_stat.max_row
    for name, count in per_file_count.items():
        ws_stat.append([name, count])

    # 數值欄位的平均/最小/最大
    numeric_stats = []
    for col_idx, header in enumerate(headers):
        values = [r[col_idx] for r in rows if isinstance(r[col_idx], (int, float))]
        if values:
            numeric_stats.append(
                (header, sum(values) / len(values), min(values), max(values))
            )
    if numeric_stats:
        ws_stat.append([])
        ws_stat.append(["數值欄位", "平均", "最小", "最大"])
        for header, avg, lo, hi in numeric_stats:
            ws_stat.append([header, round(avg, 3), lo, hi])

    for cell in ws_stat["A"]:
        cell.font = Font(bold=True)
    ws_stat.column_dimensions["A"].width = 30

    # 各檔案筆數長條圖
    chart = BarChart()
    chart.title = "各檔案資料筆數"
    data = Reference(ws_stat, min_col=2, min_row=chart_start,
                     max_row=chart_start + len(per_file_count))
    cats = Reference(ws_stat, min_col=1, min_row=chart_start + 1,
                     max_row=chart_start + len(per_file_count))
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    ws_stat.add_chart(chart, "E2")

    timestamp = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
    out_path = files[0].parent / f"合併結果_{timestamp}.xlsx"
    out.save(out_path)
    return out_path


def create_demo_files(folder: Path, log=print) -> None:
    """產生三個假的機台日報表,方便測試合併功能。"""
    random.seed()
    start = datetime.date.today() - datetime.timedelta(days=7)
    for tool_name in ["機台A", "機台B", "機台C"]:
        wb = Workbook()
        ws = wb.active
        ws.append(["日期", "機台", "批號", "產出片數", "缺陷數", "良率(%)"])
        for day in range(7):
            date = start + datetime.timedelta(days=day)
            for lot in range(random.randint(2, 4)):
                wafers = random.choice([24, 25])
                defects = random.randint(0, 30)
                yield_pct = round(100 * (1 - defects / (wafers * 100)), 2)
                ws.append([
                    date.strftime("%Y-%m-%d"),
                    tool_name,
                    f"LOT{date.strftime('%m%d')}{lot + 1:02d}",
                    wafers,
                    defects,
                    yield_pct,
                ])
        path = folder / f"日報_{tool_name}.xlsx"
        wb.save(path)
        log(f"  已建立 {path.name}")


# ---------------------------------------------------------------------------
# 圖形介面
# ---------------------------------------------------------------------------

class App(tk.Tk if tk else object):  # 無 tkinter 的環境仍可匯入上方核心函式
    def __init__(self):
        super().__init__()
        self.title(APP_TITLE)
        self.geometry("640x460")
        self.minsize(560, 400)

        frame = ttk.Frame(self, padding=12)
        frame.pack(fill="both", expand=True)

        ttk.Label(frame, text="來源資料夾(裡面放要合併的 .xlsx 檔):").pack(anchor="w")

        row = ttk.Frame(frame)
        row.pack(fill="x", pady=(4, 8))
        self.folder_var = tk.StringVar()
        ttk.Entry(row, textvariable=self.folder_var).pack(
            side="left", fill="x", expand=True)
        ttk.Button(row, text="瀏覽...", command=self.browse).pack(
            side="left", padx=(6, 0))

        buttons = ttk.Frame(frame)
        buttons.pack(fill="x", pady=(0, 8))
        self.merge_btn = ttk.Button(
            buttons, text="開始合併", command=self.start_merge)
        self.merge_btn.pack(side="left")
        ttk.Button(buttons, text="產生示範資料", command=self.make_demo).pack(
            side="left", padx=(6, 0))

        ttk.Label(frame, text="執行紀錄:").pack(anchor="w")
        self.log_box = tk.Text(frame, height=14, state="disabled")
        self.log_box.pack(fill="both", expand=True, pady=(4, 0))

        self.log(f"{APP_TITLE} 已啟動,請選擇資料夾。")

    # --- 介面輔助 ---

    def log(self, message: str) -> None:
        def append():
            self.log_box.configure(state="normal")
            self.log_box.insert("end", message + "\n")
            self.log_box.see("end")
            self.log_box.configure(state="disabled")
        self.after(0, append)

    def browse(self) -> None:
        folder = filedialog.askdirectory(title="選擇來源資料夾")
        if folder:
            self.folder_var.set(folder)

    def get_folder(self) -> Path | None:
        folder = self.folder_var.get().strip()
        if not folder:
            messagebox.showwarning(APP_TITLE, "請先選擇資料夾。")
            return None
        path = Path(folder)
        if not path.is_dir():
            messagebox.showerror(APP_TITLE, f"找不到資料夾:\n{folder}")
            return None
        return path

    # --- 按鈕動作 ---

    def make_demo(self) -> None:
        folder = self.get_folder()
        if folder is None:
            return
        self.log("產生示範資料...")
        create_demo_files(folder, log=self.log)
        self.log("完成!可以按「開始合併」測試。")

    def start_merge(self) -> None:
        folder = self.get_folder()
        if folder is None:
            return
        files = find_excel_files(folder)
        if not files:
            messagebox.showwarning(
                APP_TITLE, "這個資料夾裡沒有 .xlsx 檔案。\n"
                "可以先按「產生示範資料」建立測試檔。")
            return
        self.merge_btn.configure(state="disabled")
        threading.Thread(target=self._merge_worker, args=(files,),
                         daemon=True).start()

    def _merge_worker(self, files: list[Path]) -> None:
        try:
            self.log(f"開始合併 {len(files)} 個檔案...")
            out_path = merge_excel_files(files, log=self.log)
            self.log(f"合併完成!輸出檔案:{out_path.name}")
            self.after(0, lambda: messagebox.showinfo(
                APP_TITLE, f"合併完成!\n\n輸出檔案:\n{out_path}"))
        except Exception as exc:  # 把錯誤顯示在介面上,而不是悄悄當掉
            self.log(f"發生錯誤:{exc}")
            self.after(0, lambda: messagebox.showerror(
                APP_TITLE, f"合併失敗:\n{exc}"))
        finally:
            self.after(0, lambda: self.merge_btn.configure(state="normal"))


if __name__ == "__main__":
    if tk is None:
        raise SystemExit("此環境沒有 tkinter,無法開啟圖形介面。")
    App().mainloop()
